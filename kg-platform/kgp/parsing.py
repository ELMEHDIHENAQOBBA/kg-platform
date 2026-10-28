"""Document parsing: bytes -> (title, text). Supports txt/md/html/pdf/docx/json/eml."""
import io
import json
import re
from email import policy
from email.parser import BytesParser


def _html(data):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    for t in soup(["script", "style", "nav", "footer"]):
        t.decompose()
    return title, soup.get_text("\n", strip=True)


def _pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    meta_title = (reader.metadata.title if reader.metadata else "") or ""
    return meta_title, text


def _docx(data):
    from docx import Document
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(". ".join(c.text for c in row.cells))
    return (d.core_properties.title or ""), "\n".join(parts)


def _json(data):
    obj = json.loads(data.decode("utf-8"))
    out = []

    def walk(x):
        if isinstance(x, str):
            out.append(x)
        elif isinstance(x, dict):
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
    walk(obj)
    return "", "\n".join(out)


def _eml(data):
    msg = BytesParser(policy=policy.default).parsebytes(data)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    return str(msg["subject"] or ""), f"Subject: {msg['subject']}\n{text}"


def parse_bytes(name, data):
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else "txt"
    if ext in ("html", "htm"):
        return _html(data)
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    if ext == "json":
        return _json(data)
    if ext == "eml":
        return _eml(data)
    text = data.decode("utf-8", errors="replace")
    first = next((l.strip() for l in text.splitlines() if l.strip()), "")
    return first[:120], text


def chunk_text(text, size=1500):
    """Paragraph-aware chunking."""
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=\S)", text) if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if cur and len(cur) + len(p) > size:
            chunks.append(cur)
            cur = ""
        cur = f"{cur}\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks
