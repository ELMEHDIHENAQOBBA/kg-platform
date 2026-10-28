"""Generates the binary sample sources (PDF, DOCX, SQLite). Run: python scripts/generate_samples.py"""
import sqlite3
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
docs = ROOT / "data/sources/documents"
db_dir = ROOT / "data/sources/db"
docs.mkdir(parents=True, exist_ok=True)
db_dir.mkdir(parents=True, exist_ok=True)

PDF_LINES = [
    "Cedar Capital - Investor update",
    "Cedar Capital is headquartered in London.",
    "Thomas Weber is the Managing Partner of Cedar Capital.",
    "Cedar Capital Partners invested in Lumen Robotics in 2023.",
    "Cedar Capital invested in Atlas Dynamics in 2022.",
    "Cedar Capital invested in NovaGraph in 2022.",
]
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
c = canvas.Canvas(str(docs / "cedar_capital_investor_update.pdf"), pagesize=A4)
c.setTitle("Cedar Capital investor update")
y = 800
for i, line in enumerate(PDF_LINES):
    c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 14 if i == 0 else 11)
    c.drawString(60, y, line)
    y -= 24
c.save()

from docx import Document
d = Document()
d.core_properties.title = "Lumen Robotics company profile"
d.add_heading("Lumen Robotics company profile", 1)
d.add_paragraph("Lumen Robotics is headquartered in Berlin. Lena Fischer founded Lumen Robotics in 2016. "
                "Lena Fischer is the CEO of Lumen Robotics.")
d.add_paragraph("Lumen Robotics develops LumenArm. Lumen Robotics uses PyTorch for grasp planning. "
                "Lumen Robotics competes with Atlas Dynamics. Lumen Robotics partnered with Atlas Dynamics in 2025.")
d.save(docs / "lumen_robotics_profile.docx")

con = sqlite3.connect(db_dir / "crm.sqlite")
con.executescript("""
DROP TABLE IF EXISTS education; DROP TABLE IF EXISTS employment;
DROP TABLE IF EXISTS partnerships; DROP TABLE IF EXISTS locations;
CREATE TABLE education(person TEXT, school TEXT);
CREATE TABLE employment(person TEXT, company TEXT);
CREATE TABLE partnerships(org_a TEXT, org_b TEXT);
CREATE TABLE locations(org TEXT, city TEXT);
""")
con.executemany("INSERT INTO education VALUES (?,?)", [
    ("Amina El Fassi", "Rabat Institute of Technology"), ("Karim Benali", "RIT"),
    ("Omar Tazi", "Rabat Institute of Technology"), ("Nadia Chraibi", "Rabat Institute of Technology")])
con.executemany("INSERT INTO employment VALUES (?,?)", [("Omar Tazi", "Atlas Dynamics")])
con.executemany("INSERT INTO partnerships VALUES (?,?)", [("Orion Cloud", "Atlas Dynamics")])
con.executemany("INSERT INTO locations VALUES (?,?)", [
    ("Rabat Institute of Technology", "Rabat"), ("Orion Cloud", "Amsterdam")])
con.commit()
con.close()
print("samples generated")
